
import os,json,uuid,hmac,hashlib,time,secrets,asyncio,re
from datetime import datetime,timezone
from io import BytesIO
from pathlib import Path
from fastapi import FastAPI,Request,HTTPException,UploadFile,File
from fastapi.responses import FileResponse,JSONResponse,Response
from fastapi.staticfiles import StaticFiles
from core.models import Inputs,Generated,SaveRequest,ReviewRequest
from core.knowledge import ROOT,RUBRICS,MANIFEST,retrieve,validate,rubric_for
from core.storage import Store,valid_id
from core import ai
from core.exporters import export
app=FastAPI(title='PDP Integration Studio',docs_url=None,redoc_url=None)
store=Store()
SECRET=os.getenv('PDP_SESSION_SECRET')
if not SECRET and not os.getenv('VERCEL'):SECRET='local-development-only'
SEMAPHORE=asyncio.Semaphore(3)
RATE={}
def sign(id):return hmac.new(SECRET.encode(),id.encode(),hashlib.sha256).hexdigest()
def now():return datetime.now(timezone.utc).isoformat()
@app.middleware('http')
async def boundary(request:Request,call_next):
    if not SECRET:return JSONResponse({'detail':'Chưa cấu hình bảo vệ không gian làm việc.'},503)
    if request.method not in ['GET','HEAD','OPTIONS']:
        origin=request.headers.get('origin')
        if origin and origin.rstrip('/')!=str(request.base_url).rstrip('/'):return JSONResponse({'detail':'Yêu cầu khác nguồn bị từ chối.'},403)
    cookie=request.cookies.get('pdp_workspace','');bits=cookie.split('.')
    owner=None
    if len(bits)==2:
        try:
            valid_id(bits[0])
            if hmac.compare_digest(sign(bits[0]),bits[1]):owner=bits[0]
        except HTTPException:pass
    fresh=not owner;owner=owner or str(uuid.uuid4());request.state.owner=owner
    response=await call_next(request)
    if fresh:response.set_cookie('pdp_workspace',owner+'.'+sign(owner),httponly=True,secure=bool(os.getenv('VERCEL')),samesite='strict',max_age=365*24*3600)
    response.headers['X-Content-Type-Options']='nosniff'
    response.headers['Referrer-Policy']='same-origin'
    response.headers['Content-Security-Policy']="default-src 'self'; script-src 'self'; style-src 'self'; img-src 'self' data:; connect-src 'self'; frame-ancestors 'none'; base-uri 'self'; form-action 'self'"
    if request.url.path.startswith('/api'):response.headers['Cache-Control']='no-store'
    return response
@app.get('/')
async def index():return FileResponse(ROOT/'static/index.html')
@app.get('/api/config')
async def config():
    return {'version':'2.0','ai_ready':bool(os.getenv('OPENAI_API_KEY')),'storage':'private_blob' if store.token else 'local' if not os.getenv('VERCEL') else 'unavailable','review_ready':bool(os.getenv('PDP_REVIEW_CODE')),'rubrics':list(RUBRICS.values()),'sources':MANIFEST['documents']}
def make_doc(inputs,content,sources,parent='',kind='ai',response_id=None):
    return {'id':str(uuid.uuid4()),'parent_id':parent,'created_at':now(),'status':'draft','origin':kind,'inputs':inputs.model_dump(),'content':content.model_dump(),'rubrics':rubric_for(inputs.indicators),'sources':sources,'validation':validate(inputs,content),'review':None,'response_id':response_id}
async def reference_for(request,inputs):
    if inputs.mode!='V2':return None
    if not inputs.reference_id:raise HTTPException(422,'Chọn bản tham chiếu V1 đã duyệt trước khi tạo V2.')
    ref=await store.get(request.state.owner,inputs.reference_id)
    if ref['status']!='approved' or ref['inputs']['mode']!='V1' or ref['content']['integration']!='Tích hợp':raise HTTPException(422,'V2 cần bản V1 tích hợp đã được duyệt.')
    if inputs.lesson!=ref['inputs']['lesson'] or inputs.subject_goal!=ref['inputs']['subject_goal'] or inputs.indicators!=ref['inputs']['indicators']:raise HTTPException(422,'Bài học, YCCĐ và chỉ báo V2 phải khớp V1 đã duyệt. Hãy chọn lại bản V1.')
    return ref
@app.post('/api/generate')
async def generate(request:Request,inputs:Inputs):
    if any(c not in RUBRICS for c in inputs.indicators):raise HTTPException(422,'Mã PDP không hợp lệ.')
    ref=await reference_for(request,inputs)
    owner=request.state.owner;stamp=time.monotonic()
    if stamp-RATE.get(owner,-60)<20:raise HTTPException(429,'Vui lòng chờ 20 giây trước lượt tạo tiếp theo.')
    RATE[owner]=stamp
    if len(RATE)>10000:
        for k,v in list(RATE.items()):
            if stamp-v>3600:RATE.pop(k,None)
    async with SEMAPHORE:content,sources,response_id=await ai.generate(inputs,ref)
    doc=make_doc(inputs,content,sources,response_id=response_id)
    await store.put(owner,doc)
    return doc
@app.post('/api/save')
async def save(request:Request,data:SaveRequest):
    if any(c not in RUBRICS for c in data.inputs.indicators):raise HTTPException(422,'Mã PDP không hợp lệ.')
    parent=await store.get(request.state.owner,data.parent_id) if data.parent_id else None
    await reference_for(request,data.inputs)
    doc=make_doc(data.inputs,data.content,retrieve(data.inputs),data.parent_id,kind=parent['origin'] if parent else 'manual')
    await store.put(request.state.owner,doc);return doc
@app.get('/api/documents')
async def documents(request:Request):
    docs=await store.list(request.state.owner)
    return [{'id':d['id'],'parent_id':d['parent_id'],'title':d['content']['title'],'mode':d['inputs']['mode'],'status':d['status'],'created_at':d['created_at'],'lesson':d['inputs']['lesson'],'inputs':d['inputs'],'origin':d['origin']} for d in docs]
@app.get('/api/documents/{id}')
async def document(request:Request,id:str):return await store.get(request.state.owner,id)
@app.post('/api/documents/{id}/approve')
async def approve(request:Request,id:str,data:ReviewRequest):
    code=os.getenv('PDP_REVIEW_CODE')
    if not code:raise HTTPException(503,'Chưa cấu hình quyền duyệt.')
    owner=request.state.owner;ratekey='review-'+owner;stamp=time.monotonic()
    if stamp-RATE.get(ratekey,-3)<3:raise HTTPException(429,'Vui lòng chờ trước khi thử mã duyệt tiếp.')
    RATE[ratekey]=stamp
    if not hmac.compare_digest(data.code,code):raise HTTPException(403,'Mã duyệt chưa đúng.')
    doc=await store.get(owner,id)
    if doc['origin']=='sample':raise HTTPException(422,'Ví dụ minh họa không được duyệt chính thức. Hãy tạo bản mới từ yêu cầu của bạn.')
    checks=validate(Inputs.model_validate(doc['inputs']),Generated.model_validate(doc['content']))
    if checks['errors']:raise HTTPException(422,'Cần xử lý các mục thiếu cấu trúc trước khi duyệt.')
    if not data.source_confirmed:raise HTTPException(422,'Người duyệt cần xác nhận đã đối chiếu YCCĐ và tính đúng của Toán.')
    doc['parent_id']=doc['id'];doc['id']=str(uuid.uuid4());doc['created_at']=now();doc['status']='approved';doc['validation']=checks
    doc['review']={'reviewer':data.reviewer,'note':data.note,'time':now(),'source_confirmed':True,'method':'Mã duyệt của người quản lý; tên do người duyệt khai báo'}
    await store.put(owner,doc);return doc
@app.get('/api/documents/{id}/export/{kind}')
async def download(request:Request,id:str,kind:str):
    if kind not in ['docx','xlsx','pdf']:raise HTTPException(404,'Định dạng không hỗ trợ.')
    doc=await store.get(request.state.owner,id)
    result=await asyncio.to_thread(export,doc,kind)
    mime={'docx':'application/vnd.openxmlformats-officedocument.wordprocessingml.document','xlsx':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','pdf':'application/pdf'}[kind]
    return Response(result,media_type=mime,headers={'Content-Disposition':f'attachment; filename="PDP-{doc["inputs"]["mode"]}-{id[:8]}.{kind}"'})
@app.post('/api/extract')
async def extract(file:UploadFile=File(...)):
    raw=await file.read(3_000_001)
    if len(raw)>3_000_000:raise HTTPException(413,'Tệp tối đa 3 MB. Hãy tách trang/phần cần dùng.')
    suffix=Path(file.filename or '').suffix.lower()
    try:
        if suffix=='.pdf':
            from pypdf import PdfReader
            reader=PdfReader(BytesIO(raw))
            if len(reader.pages)>80:raise HTTPException(413,'Tệp tối đa 80 trang. Hãy tách phần cần dùng.')
            text='\n'.join(f'[Trang {i+1}] '+(p.extract_text() or '') for i,p in enumerate(reader.pages))
        elif suffix=='.docx':
            from docx import Document
            d=Document(BytesIO(raw));text='\n'.join([p.text for p in d.paragraphs]+[' | '.join(c.text for c in row.cells) for table in d.tables for row in table.rows])
        elif suffix=='.xlsx':
            from openpyxl import load_workbook
            w=load_workbook(BytesIO(raw),read_only=True,data_only=True)
            text='\n'.join(f'[{s.title} hàng {i}] '+' | '.join(str(c or '') for c in row) for s in w for i,row in enumerate(s.iter_rows(values_only=True),1) if i<=500)
        elif suffix=='.txt':text=raw.decode('utf8')
        else:raise HTTPException(422,'Hỗ trợ PDF, DOCX, XLSX và TXT.')
    except HTTPException:raise
    except Exception:raise HTTPException(422,'Không đọc được tệp. Hãy thử dán đoạn nguồn trực tiếp.')
    if len(re.sub(r'\[Trang \d+\]','',text).strip())<30:raise HTTPException(422,'Tệp chưa có văn bản đọc được. Với PDF scan, cần OCR hoặc dán phần YCCĐ.')
    return {'name':file.filename,'text':text[:24000],'truncated':len(text)>24000}
@app.get('/api/sample')
async def sample(request:Request):
    data=json.loads((ROOT/'data/sample.json').read_text(encoding='utf8'))
    doc=make_doc(Inputs.model_validate(data['inputs']),Generated.model_validate(data['content']),[],kind='sample')
    await store.put(request.state.owner,doc)
    return doc
app.mount('/static',StaticFiles(directory=ROOT/'static'),name='static')
