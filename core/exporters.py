
import io,base64
from xml.sax.saxutils import escape
from .knowledge import ROOT
LEVELS=['Khởi đầu / Định hình','Phát triển','Thành thạo','Nâng cao']
COLUMNS=['Chủ đề','Bài học','YCCĐ môn học','Mã PDP','Điểm chạm','Phương thức sư phạm','Minh chứng','Đánh giá']
def blocks(doc):
    c=doc['content'];i=doc['inputs']
    result=[('PDP INTEGRATION STUDIO · '+i['mode'],c['title']),('Thông tin',f"Toán lớp 5 · Kết nối tri thức · {i['minutes']} phút\nNgười soạn: {i['author']}\nTrạng thái: "+('Đã duyệt' if doc['status']=='approved' else 'Bản nháp · Cần giáo viên rà soát')),('Yêu cầu cần đạt môn học',i['subject_goal']),('Quyết định tích hợp',c['integration']+'\n'+c['rationale'])]
    for s in c['sections']:result.append((s['heading'],s['content']))
    if c['reference_rows']:
        for row in c['reference_rows']:
            result.append(('Bảng tham chiếu · '+row['lesson'],'\n'.join(label+': '+(', '.join(value) if isinstance(value,list) else value) for label,value in zip(COLUMNS,row.values()))))
    for e in c['evidence']:result.append(('Minh chứng · '+e['kind'],e['description']+'\nThu thập: '+e['method']))
    for key,title in [('reflection_before','Trước khi làm'),('reflection_during','Trong khi làm'),('reflection_after','Sau khi làm')]:result.append(('Nhật ký · '+title,c[key]))
    for r in doc['rubrics']:
        result.append((r['code']+' · '+r['name'],'\n\n'.join(level+': '+text for level,text in zip(LEVELS,r['levels']))+'\nNguồn: '+r['source']))
    result.extend([('Nguồn YCCĐ do giáo viên bổ sung',i['source_name']+'\n'+i['source_excerpt']),('Nguồn khung đã sử dụng','\n'.join(s['filename']+' · '+s['locator'] for s in doc['sources'])),('Lưu ý','\n'.join(doc['validation']['warnings']+doc['validation']['errors']))])
    if doc.get('review'):result.append(('Dấu vết duyệt',doc['review']['reviewer']+' · '+doc['review']['time']+'\n'+doc['review']['note']))
    return result
def export(doc,kind):
    output=io.BytesIO()
    if kind=='docx':
        from docx import Document
        from docx.shared import Pt
        d=Document();d.styles['Normal'].font.name='Arial';d.styles['Normal'].font.size=Pt(11)
        for n,(heading,text) in enumerate(blocks(doc)):
            d.add_heading(heading,0 if n==0 else 1)
            for line in text.split('\n'):d.add_paragraph(line)
        d.save(output)
    elif kind=='xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font,Alignment,PatternFill
        w=Workbook();s=w.active;s.title='Nội dung';s.append(['Mục','Nội dung'])
        for h,t in blocks(doc):s.append([h,t])
        s.column_dimensions['A'].width=35;s.column_dimensions['B'].width=100
        if doc['content']['reference_rows']:
            t=w.create_sheet('Bảng tham chiếu');t.append(COLUMNS)
            for r in doc['content']['reference_rows']:t.append([', '.join(v) if isinstance(v,list) else v for v in r.values()])
            for col in 'ABCDEFGH':t.column_dimensions[col].width=35
        r=w.create_sheet('Rubric gốc');r.append(['Mã','Chỉ báo']+LEVELS+['Nguồn'])
        for item in doc['rubrics']:r.append([item['code'],item['name']]+item['levels']+[item['source']])
        for col in 'ABCDEFG':r.column_dimensions[col].width=40
        for sheet in w:
            sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
            for cell in sheet[1]:cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='172E35')
            for row in sheet.iter_rows(min_row=2):
                sheet.row_dimensions[row[0].row].height=min(300,max(45,max(len(str(c.value or '')) for c in row)//3))
                for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
        w.save(output)
    elif kind=='pdf':
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer
        from reportlab.lib.styles import ParagraphStyle
        font=io.BytesIO(base64.b64decode((ROOT/'data/font.b64').read_text()))
        if 'PDPUnicode' not in pdfmetrics.getRegisteredFontNames():pdfmetrics.registerFont(TTFont('PDPUnicode',font))
        style=ParagraphStyle('body',fontName='PDPUnicode',fontSize=10,leading=15,spaceAfter=7)
        heading=ParagraphStyle('heading',parent=style,fontSize=14,leading=20,spaceBefore=14)
        flow=[]
        for h,t in blocks(doc):
            flow.append(Paragraph(escape(h),heading))
            for line in t.split('\n'):flow.append(Paragraph(escape(line) or ' ',style))
            flow.append(Spacer(1,5))
        SimpleDocTemplate(output,title=doc['content']['title'],author=doc['inputs']['author'],leftMargin=42,rightMargin=42,topMargin=40,bottomMargin=40).build(flow)
    else:raise ValueError('Unsupported export')
    return output.getvalue()
